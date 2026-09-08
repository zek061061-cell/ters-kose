"""Import Riskbudur-style Excel/CSV history into Ters Köşe v7.

Expected headers follow the public Riskbudur "Veri Güncelle" schema:
Ülke, Lig, Sezon, Gün, Tarih, Yıl, Saat, Hafta, Hakem,
Ev Sahibi, Deplasman, MS, İY, İY S, MS S, İY-MS, +6,
2.5, 3.5, KG, İY 0.5, İY 1.5.

The importer never invents scores. It keeps only rows whose country/league
resolve to the focused 57-league catalog.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

DATA_DIR = Path(__file__).with_name("data")
CATALOG_PATH = DATA_DIR / "riskbudur_57_catalog.json"
DEFAULT_OUTPUT = DATA_DIR / "riskbudur_57_history.json"

HEADER_ALIASES = {
    "ulke": "country", "ülke": "country",
    "lig": "league", "sezon": "season", "gun": "day", "gün": "day",
    "tarih": "date", "yil": "year", "yıl": "year", "saat": "time",
    "hafta": "week", "hakem": "referee",
    "evsahibi": "home", "ev sahibi": "home",
    "deplasman": "away",
    "ms": "ft_result", "iy": "ht_result",
    "iys": "ht_score", "iy s": "ht_score",
    "mss": "ft_score", "ms s": "ft_score",
    "iy-ms": "ht_ft", "iyms": "ht_ft",
    "+6": "six_plus", "2.5": "over25", "3.5": "over35",
    "kg": "btts", "iy0.5": "ht_over05", "iy 0.5": "ht_over05",
    "iy1.5": "ht_over15", "iy 1.5": "ht_over15",
}


def _key(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").strip().casefold())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9+.-]+", "", text)


def _header(value: object) -> str:
    raw = str(value or "").strip().casefold()
    normalized = unicodedata.normalize("NFKD", raw)
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    spaced = re.sub(r"\s+", " ", normalized).strip()
    compact = spaced.replace(" ", "")
    return HEADER_ALIASES.get(spaced) or HEADER_ALIASES.get(compact) or compact


def _score(value: object):
    text = str(value or "").strip()
    m = re.search(r"(\d+)\s*[-:]\s*(\d+)", text)
    return (int(m.group(1)), int(m.group(2))) if m else (None, None)


def _iso_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def _truth(value: object) -> bool | None:
    text = str(value or "").strip().casefold()
    if not text:
        return None
    if text in {"1", "evet", "var", "yes", "true", "üst", "ust"}:
        return True
    if text in {"0", "hayır", "hayir", "yok", "no", "false", "alt"}:
        return False
    return None


def _catalog_index() -> dict[tuple[str, str], dict]:
    payload = json.loads(CATALOG_PATH.read_text(encoding="utf-8"))
    index = {}
    for league in payload["leagues"]:
        countries = [league["country"]]
        names = [league["canonical_name"], *league.get("aliases", [])]
        for country in countries:
            for name in names:
                index[(_key(country), _key(name))] = league
    return index


def _read_rows(path: Path):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            for row in reader:
                yield {_header(k): v for k, v in row.items()}
        return

    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [_header(v) for v in next(rows)]
        except StopIteration:
            return
        for values in rows:
            yield {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
        return

    raise ValueError("Only .xlsx, .xlsm and .csv are supported")


def import_file(path: Path) -> dict:
    league_index = _catalog_index()
    output, rejected = [], []
    seen = set()

    for n, raw in enumerate(_read_rows(path), start=2):
        country = str(raw.get("country") or "").strip()
        league_name = str(raw.get("league") or "").strip()
        league = league_index.get((_key(country), _key(league_name)))
        if not league:
            rejected.append({"row": n, "reason": "league_not_in_57", "country": country, "league": league_name})
            continue

        match_date = _iso_date(raw.get("date"))
        home = str(raw.get("home") or "").strip()
        away = str(raw.get("away") or "").strip()
        if not match_date or not home or not away:
            rejected.append({"row": n, "reason": "missing_date_or_team"})
            continue

        ht_home, ht_away = _score(raw.get("ht_score"))
        ft_home, ft_away = _score(raw.get("ft_score"))
        completed = ft_home is not None and ft_away is not None
        identity = (league["league_id"], match_date, _key(home), _key(away))
        if identity in seen:
            continue
        seen.add(identity)

        total = (ft_home + ft_away) if completed else None
        btts = (ft_home > 0 and ft_away > 0) if completed else _truth(raw.get("btts"))
        htft = str(raw.get("ht_ft") or "").strip()
        if not htft and ht_home is not None and ft_home is not None:
            ht_res = "1" if ht_home > ht_away else "2" if ht_away > ht_home else "X"
            ft_res = "1" if ft_home > ft_away else "2" if ft_away > ft_home else "X"
            htft = f"{ht_res}/{ft_res}"

        kickoff = str(raw.get("time") or "").strip()
        if kickoff and re.fullmatch(r"\d{1,2}:\d{2}", kickoff):
            kickoff = f"{match_date}T{kickoff}:00"

        output.append({
            "event_id": f"RB57|{league['league_id']}|{match_date}|{_key(home)}|{_key(away)}",
            "date": match_date,
            "kickoff": kickoff or match_date,
            "league_id": league["league_id"],
            "league": f"{league['country']} {league['canonical_name']}",
            "country": league["country"],
            "season": str(raw.get("season") or ""),
            "week": raw.get("week") if raw.get("week") not in ("", None) else None,
            "referee": str(raw.get("referee") or "").strip(),
            "home": home,
            "away": away,
            "completed": completed,
            "ht_home": ht_home, "ht_away": ht_away,
            "ft_home": ft_home, "ft_away": ft_away,
            "ht_ft": htft,
            "six_plus": (total >= 6) if total is not None else _truth(raw.get("six_plus")),
            "over25": (total >= 3) if total is not None else _truth(raw.get("over25")),
            "over35": (total >= 4) if total is not None else _truth(raw.get("over35")),
            "btts": btts,
            "ht_over05": ((ht_home + ht_away) >= 1) if ht_home is not None else _truth(raw.get("ht_over05")),
            "ht_over15": ((ht_home + ht_away) >= 2) if ht_home is not None else _truth(raw.get("ht_over15")),
            "source": "Riskbudur-format import",
        })

    output.sort(key=lambda row: (row["date"], row["league_id"], row["home"], row["away"]))
    return {
        "schema_version": 1,
        "profile": "riskbudur-57",
        "source_file": path.name,
        "league_count": len({row["league_id"] for row in output}),
        "match_count": len(output),
        "rejected_count": len(rejected),
        "matches": output,
        "rejected": rejected[:500],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    args = parser.parse_args()
    result = import_file(Path(args.input))
    Path(args.output).write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: result[k] for k in ("league_count", "match_count", "rejected_count")}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
